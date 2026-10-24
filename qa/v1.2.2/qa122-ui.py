from ui12 import *
from pptx import Presentation
import ui12,sys
ROOT=Path('work/qa122/ui');ROOT.mkdir(parents=True,exist_ok=True)
def robust_adb(*args):
 if args[0]=='shell':args=('shell','-T','-x',*args[1:])
 return subprocess.run([ADB,'-s','emulator-5554',*args],capture_output=True,check=True,timeout=60).stdout
ui12.adb=robust_adb;adb=robust_adb

def menu(label):tap('Действия с документом');time.sleep(.5);tap(label);time.sleep(.5)
def meta(name):
 xml=ET.fromstring(adb('exec-out','run-as','app.quietread','cat','shared_prefs/library.xml'))
 return next(i for i in json.loads(next(n.text for n in xml if n.get('name')=='items')) if i['name']==name)
def point(name,t):
 node=next(n for n in t.iter('node') if n.get('content-desc','').startswith('Документ.'));b=bounds(node)
 log=adb('logcat','-d','-s','QuietRead').decode(errors='replace');typ,parts,w,h=map(int,re.findall(r'OPEN_OK '+re.escape(name)+r' type=(-?\d+) parts=(\d+) size=(\d+)x(\d+)',log)[-1]);state=meta(name);scale=state['zoom']
 if typ==1:x,y=500,150
 elif typ==2:
  shape=Presentation('outputs/QuietRead/app/src/debug/assets/qa/sample.pptx').slides[0].shapes.title;x=(shape.left+shape.width/2)/635;y=(shape.top+shape.height/2)/635
 elif typ==3:x,y=3*1440/2.54,2.3*1440/2.54
 elif typ==-1:x,y=2000,1500
 else:x,y=2200,2000
 return typ,round(b[0]+max(32,(b[2]-b[0]-w*scale)/2)-state['x']*scale+x*scale),round(b[1]+32-state['y']*scale+y*scale)
def run(name):
 open_doc(name);menu('По ширине / сбросить масштаб');menu('Редактировать');t=tree();typ,x,y=point(name,t);adb('shell','input','tap',str(x),str(y));time.sleep(.9);assert keyboard(),name+' no IME'
 t=tree();assert any(n.get('content-desc','').startswith('Документ.') for n in t.iter('node')),name+' canvas absent'
 assert not any(n.get('text') in ['Изменить текст','Заменить строку PDF','Новая надпись'] for n in t.iter('node')),name+' modal editor'
 field=next(n for n in t.iter('node') if n.get('content-desc')==('Значение или формула ячейки' if typ==1 else 'Текст в документе'))
 if typ!=1:
  canvas=next(n for n in t.iter('node') if n.get('content-desc','').startswith('Документ.'));b=bounds(canvas);assert b[3]-b[1]>=105,(name,'No visible document',b)
 (ROOT/(name+'.png')).write_bytes(adb('exec-out','screencap','-p'))
 marker='U122_'+name.rsplit('.',1)[1].upper()+'_'+str(int(time.time()))
 apply=next(n for n in t.iter('node') if n.get('content-desc')==('Применить значение ячейки' if typ==1 else 'Применить текст'))
 adb('shell','input','tap',*map(str,center(field)));adb('shell','input','keycombination','113','29');adb('shell','input','text',marker);adb('shell','input','tap',*map(str,center(apply)))
 data=draft(name)
 if typ==-1:
  from pypdf import PdfReader
  content=PdfReader(io.BytesIO(data)).pages[0].extract_text()
 else:content=ZipFile(io.BytesIO(data)).read('content.xml').decode('utf8')
 assert marker in content,(name,'not saved');assert not keyboard(),name+' IME still open'
 print(json.dumps({'file':name,'inline_editor':True,'document_visible':True,'automatic_ime':True,'saved':True},ensure_ascii=False),flush=True)
if __name__=='__main__':
 for name in sys.argv[1:] or ['sample.docx','sample.odt','sample.pptx','sample.odg','sample.xlsx','sample.ods','pdf-ui12.pdf']:run(name)
